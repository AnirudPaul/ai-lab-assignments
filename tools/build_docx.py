"""Render the executed notebook as a Word document.

    python tools/build_docx.py

Produces AI_Lab_Assignment.docx with:
  * a decorative page border on every page,
  * a running header carrying the course code and title,
  * the VIT title page,
  * markdown rendered as real Word headings / tables / lists,
  * code in shaded monospace blocks and console output in bordered blocks,
  * every figure embedded at full width.

The notebook is the single source of truth -- nothing is re-run here, so the document
always shows exactly what the stored outputs contain.
"""
from __future__ import annotations

import base64
import io
import re
import sys
from pathlib import Path

import nbformat
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent

COURSE_CODE = "CSD5004"
COURSE_NAME = "Artificial Intelligence and Machine Learning"
STUDENT = "Anirud Paul"
REGNO = "26MCF10001"
CLASSNBR = "BL2026270100790"
SLOT = "B11+B12+B13"
ROOM = "AB02-202"
FACULTY = "URIKHIMBAM BOBY CLINTON"

BRAND = RGBColor(0x2A, 0x4B, 0x9B)
BRAND_HEX = "2A4B9B"
CODE_BG = "F2F4F8"
OUT_BG = "FAFAFA"
MONO = "Consolas"
BODY = "Calibri"
CONTENT_WIDTH_IN = 6.3


# --------------------------------------------------------------------------- helpers
def shade(paragraph, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(shd)


def left_bar(paragraph, color: str, size: int = 18) -> None:
    """A coloured rule down the left edge of a block."""
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bar = OxmlElement("w:left")
    bar.set(qn("w:val"), "single")
    bar.set(qn("w:sz"), str(size))
    bar.set(qn("w:space"), "6")
    bar.set(qn("w:color"), color)
    borders.append(bar)
    pPr.append(borders)


def tight(paragraph, before: int = 0, after: int = 0, line: float | None = None) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


def page_borders(section) -> None:
    """Draw a box around the whole page -- measured from the page edge, not the text."""
    sectPr = section._sectPr
    pg = OxmlElement("w:pgBorders")
    pg.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "18")          # 18 eighths of a point = 2.25pt
        el.set(qn("w:space"), "24")       # distance from page edge, in points
        el.set(qn("w:color"), BRAND_HEX)
        pg.append(el)
    sectPr.append(pg)


def field(paragraph, code: str):
    """Insert a Word field (used for PAGE / NUMPAGES)."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    return run


# ----------------------------------------------------------------- LaTeX -> Unicode
GREEK = {
    "alpha": "α", "beta": "β", "gamma": "γ", "delta": "δ", "epsilon": "ε", "zeta": "ζ",
    "eta": "η", "theta": "θ", "lambda": "λ", "mu": "μ", "nu": "ν", "xi": "ξ", "pi": "π",
    "rho": "ρ", "sigma": "σ", "tau": "τ", "phi": "φ", "chi": "χ", "psi": "ψ",
    "omega": "ω", "Gamma": "Γ", "Delta": "Δ", "Theta": "Θ", "Lambda": "Λ", "Xi": "Ξ",
    "Pi": "Π", "Sigma": "Σ", "Phi": "Φ", "Psi": "Ψ", "Omega": "Ω",
}
SYMBOLS = {
    r"\times": "×", r"\cdot": "·", r"\le": "≤", r"\leq": "≤", r"\ge": "≥", r"\geq": "≥",
    r"\neq": "≠", r"\approx": "≈", r"\pm": "±", r"\to": "→", r"\rightarrow": "→",
    r"\Rightarrow": "⇒", r"\in": "∈", r"\notin": "∉", r"\subset": "⊂", r"\infty": "∞",
    r"\sum": "Σ", r"\prod": "∏", r"\partial": "∂", r"\nabla": "∇", r"\odot": "⊙",
    r"\sqrt": "√", r"\ldots": "…", r"\dots": "…", r"\quad": "  ", r"\qquad": "    ",
    r"\,": " ", r"\;": " ", r"\!": "", r"\left": "", r"\right": "", r"\big": "",
    r"\Big": "", r"\lVert": "‖", r"\rVert": "‖", r"\lvert": "|", r"\rvert": "|",
    r"\mathbb{R}": "ℝ", r"\varepsilon": "ε", r"\sigma": "σ",
}
SUBS = str.maketrans("0123456789+-=()aeoxhijklmnprstuv",
                     "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎ₐₑₒₓₕᵢⱼₖₗₘₙₚᵣₛₜᵤᵥ")
SUPS = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")


def latex_to_text(s: str) -> str:
    """Best-effort LaTeX -> readable Unicode. Word has no TeX engine; this keeps the
    formulae legible rather than leaving raw backslashes in the document."""
    s = re.sub(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", s)
    s = re.sub(r"\\text(?:rm|bf|it)?\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\mathbf\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\operatorname\{([^{}]*)\}", r"\1", s)
    for k, v in SYMBOLS.items():
        s = s.replace(k, v)
    s = re.sub(r"\\([A-Za-z]+)", lambda m: GREEK.get(m.group(1), m.group(1)), s)
    s = re.sub(r"_\{([^{}]+)\}", lambda m: m.group(1).translate(SUBS), s)
    s = re.sub(r"\^\{([^{}]+)\}", lambda m: m.group(1).translate(SUPS), s)
    s = re.sub(r"_([A-Za-z0-9])", lambda m: m.group(1).translate(SUBS), s)
    s = re.sub(r"\^([A-Za-z0-9])", lambda m: m.group(1).translate(SUPS), s)
    return s.replace("{", "").replace("}", "").replace("$", "")


# ------------------------------------------------------------------ inline markdown
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`]+?`|\$[^$\n]+?\$)", re.S)


def add_inline(paragraph, text: str, size: float = 10.5, base_italic: bool = False):
    """Render **bold**, *italic*, `code` and $math$ into runs on a paragraph."""
    text = text.replace("&nbsp;", " ").replace("&middot;", "·").replace("&mdash;", "—")
    for part in INLINE.split(text):
        if not part:
            continue
        run = paragraph.add_run()
        run.font.size = Pt(size)
        run.italic = base_italic
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run.text = part[1:-1]
            run.font.name = MONO
            run.font.size = Pt(size - 1)
            run.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
        elif part.startswith("$") and part.endswith("$") and len(part) > 2:
            run.text = latex_to_text(part[1:-1])
            run.italic = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run.text = part[1:-1]
            run.italic = True
        else:
            run.text = part
        run.font.name = run.font.name or BODY
    return paragraph


# --------------------------------------------------------------------- block render
class DocBuilder:
    def __init__(self, doc: Document):
        self.doc = doc

    # ---- markdown --------------------------------------------------------
    def markdown(self, md: str) -> None:
        lines = md.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i]
            stripped = line.strip()

            if not stripped or stripped == "---":
                i += 1
                continue

            # display math  $$ ... $$
            if stripped.startswith("$$"):
                buf, i = self._collect_math(lines, i)
                p = self.doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                tight(p, 6, 6)
                r = p.add_run(latex_to_text(buf))
                r.italic = True
                r.font.size = Pt(11)
                r.font.name = "Cambria Math"
                continue

            # fenced code
            if stripped.startswith("```"):
                buf, i = self._collect_fence(lines, i)
                self.code_block(buf, lang_note=False)
                continue

            # table
            if stripped.startswith("|") and i + 1 < len(lines) and \
                    re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
                rows, i = self._collect_table(lines, i)
                self.table(rows)
                continue

            # headings
            if stripped.startswith("#"):
                level = len(stripped) - len(stripped.lstrip("#"))
                self.heading(stripped[level:].strip(), level)
                i += 1
                continue

            # bullet / numbered list
            m_b = re.match(r"^(\s*)[*-]\s+(.*)$", line)
            m_n = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
            if m_b or m_n:
                indent = len((m_b or m_n).group(1))
                body = m_b.group(2) if m_b else m_n.group(3)
                cont = []
                j = i + 1
                while j < len(lines) and lines[j].strip() and \
                        not re.match(r"^\s*([*-]|\d+\.)\s+", lines[j]) and \
                        not lines[j].strip().startswith(("#", "|", "```", "$$")) and \
                        lines[j].startswith(" "):
                    cont.append(lines[j].strip())
                    j += 1
                text = " ".join([body] + cont)
                style = "List Bullet" if m_b else "List Number"
                p = self.doc.add_paragraph(style=style)
                p.paragraph_format.left_indent = Inches(0.25 + 0.25 * (indent // 2))
                tight(p, 1, 3)
                add_inline(p, text)
                i = j
                continue

            # plain paragraph (join wrapped lines)
            buf = [stripped]
            j = i + 1
            while j < len(lines) and lines[j].strip() and \
                    not re.match(r"^\s*([*-]|\d+\.)\s+", lines[j]) and \
                    not lines[j].strip().startswith(("#", "|", "```", "$$", "---")):
                buf.append(lines[j].strip())
                j += 1
            p = self.doc.add_paragraph()
            tight(p, 2, 6, 1.15)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_inline(p, " ".join(buf))
            i = j

    def _collect_math(self, lines, i):
        first = lines[i].strip()
        if first.endswith("$$") and len(first) > 4:
            return first.strip("$"), i + 1
        buf, i = [first.lstrip("$")], i + 1
        while i < len(lines) and not lines[i].strip().endswith("$$"):
            buf.append(lines[i].strip())
            i += 1
        if i < len(lines):
            buf.append(lines[i].strip().rstrip("$"))
            i += 1
        return " ".join(b for b in buf if b), i

    def _collect_fence(self, lines, i):
        i += 1
        buf = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            buf.append(lines[i])
            i += 1
        return "\n".join(buf), i + 1

    def _collect_table(self, lines, i):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            raw = lines[i].strip().strip("|")
            if not re.match(r"^[\s:\-|]+$", raw):
                rows.append([c.strip() for c in raw.split("|")])
            i += 1
        return rows, i

    # ---- blocks ----------------------------------------------------------
    def heading(self, text: str, level: int) -> None:
        text = re.sub(r"^Experiment\s+", "Experiment ", text).replace("\\*", "*")
        if level == 1:
            self.doc.add_page_break()
        p = self.doc.add_paragraph()
        tight(p, 10 if level > 1 else 0, 6)
        sizes = {1: 17, 2: 14, 3: 12, 4: 11}
        run = p.add_run(re.sub(r"[*`$]", "", text))
        run.bold = True
        run.font.size = Pt(sizes.get(level, 11))
        run.font.color.rgb = BRAND
        run.font.name = BODY
        if level <= 2:
            pPr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "3")
            bottom.set(qn("w:color"), "C9D2DE")
            bdr.append(bottom)
            pPr.append(bdr)

    def code_block(self, source: str, lang_note: bool = True) -> None:
        if lang_note:
            cap = self.doc.add_paragraph()
            tight(cap, 8, 0)
            r = cap.add_run("Code")
            r.bold = True
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        for raw in source.rstrip("\n").split("\n"):
            p = self.doc.add_paragraph()
            tight(p, 0, 0, 1.0)
            p.paragraph_format.left_indent = Inches(0.08)
            shade(p, CODE_BG)
            left_bar(p, BRAND_HEX)
            run = p.add_run(raw if raw.strip() else " ")
            run.font.name = MONO
            run.font.size = Pt(8)
            rPr = run._r.get_or_add_rPr()
            rf = OxmlElement("w:rFonts")
            rf.set(qn("w:ascii"), MONO)
            rf.set(qn("w:hAnsi"), MONO)
            rPr.append(rf)

    def output_text(self, text: str, is_error: bool = False) -> None:
        text = text.rstrip("\n")
        if not text.strip():
            return
        lines = text.split("\n")
        if len(lines) > 120:                       # keep the document readable
            lines = lines[:60] + [f"... [{len(lines) - 90} lines omitted] ..."] + lines[-30:]
        for raw in lines:
            p = self.doc.add_paragraph()
            tight(p, 0, 0, 1.0)
            p.paragraph_format.left_indent = Inches(0.08)
            shade(p, OUT_BG)
            left_bar(p, "C44E52" if is_error else "9AA8BA", 12)
            run = p.add_run(raw if raw.strip() else " ")
            run.font.name = MONO
            run.font.size = Pt(7.5)
            run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
            rPr = run._r.get_or_add_rPr()
            rf = OxmlElement("w:rFonts")
            rf.set(qn("w:ascii"), MONO)
            rf.set(qn("w:hAnsi"), MONO)
            rPr.append(rf)

    def image(self, png: bytes) -> None:
        from PIL import Image

        stream = io.BytesIO(png)
        with Image.open(io.BytesIO(png)) as im:
            w, h = im.size
        width = min(CONTENT_WIDTH_IN, CONTENT_WIDTH_IN)
        if h / w * width > 8.2:                    # keep tall figures on one page
            width = 8.2 * w / h
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tight(p, 8, 8)
        p.add_run().add_picture(stream, width=Inches(width))

    def table(self, rows: list[list[str]]) -> None:
        if not rows:
            return
        ncols = max(len(r) for r in rows)
        t = self.doc.add_table(rows=0, cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci in range(ncols):
                cell = cells[ci]
                cell.text = ""
                p = cell.paragraphs[0]
                tight(p, 1, 1, 1.0)
                txt = row[ci] if ci < len(row) else ""
                add_inline(p, txt.replace("\\|", "|"), size=8.5)
                if ri == 0:
                    shade(p, "EEF1F6")
                    for r in p.runs:
                        r.bold = True
        tight(self.doc.add_paragraph(), 0, 4)


# ------------------------------------------------------------------------ title page
def title_page(doc: Document) -> None:
    logo = ROOT / "assets" / "vit_logo.png"

    sp = doc.add_paragraph()
    tight(sp, 26, 0)

    if logo.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tight(p, 0, 14)
        p.add_run().add_picture(str(logo), width=Inches(2.6))

    rule = doc.add_paragraph()
    rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(rule, 10, 10)
    pPr = rule._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    for edge in ("top", "bottom"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "12")
        el.set(qn("w:space"), "6")
        el.set(qn("w:color"), "111111")
        bdr.append(el)
    pPr.append(bdr)
    r = rule.add_run("LABORATORY RECORD")
    r.bold = True
    r.font.size = Pt(26)
    r.font.name = "Times New Roman"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, 12, 2)
    r = p.add_run(f"{COURSE_CODE}   {COURSE_NAME}")
    r.bold = True
    r.font.size = Pt(15)
    r.font.name = "Times New Roman"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, 0, 2)
    r = p.add_run("Programme Core  ·  Regular  ·  L-T-P-J-C: 2-1-1-0-4.0")
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    r.font.name = "Times New Roman"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(p, 0, 18)
    r = p.add_run("(Lecture and Tutorial, practical hours only)")
    r.italic = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    r.font.name = "Times New Roman"

    for label, value in [
        ("Student Name", STUDENT), ("Student Regn. No.", REGNO),
        ("Year", "2026-27"), ("Semester", "Fall Semester 26-27"),
        ("ClassNbr", CLASSNBR), ("Slot", SLOT), ("Class Room", ROOM),
        ("Faculty Name", FACULTY),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tight(p, 0, 6)
        a = p.add_run(f"{label}: ")
        a.bold = True
        a.font.size = Pt(12.5)
        a.font.name = "Times New Roman"
        b = p.add_run(value)
        b.font.size = Pt(12.5)
        b.font.name = "Times New Roman"

    for text, sz in [
        ("School of Computing Science Engineering and Artificial Intelligence (SCAI)", 11.5),
        ("VIT Bhopal University", 11.5),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tight(p, 10 if sz == 11.5 and "School" in text else 0, 2)
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(sz)
        r.font.name = "Times New Roman"


# ------------------------------------------------------------------------------ main
def main() -> int:
    nb_path = ROOT / "AI_Lab_Assignment.ipynb"
    out = ROOT / "AI_Lab_Assignment.docx"
    if not nb_path.exists():
        raise SystemExit(f"{nb_path} not found -- build the notebook first")

    nb = nbformat.read(nb_path, as_version=4)
    doc = Document()

    # page setup + borders
    section = doc.sections[0]
    section.page_width = Inches(8.27)      # A4
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.4)
    page_borders(section)

    normal = doc.styles["Normal"]
    normal.font.name = BODY
    normal.font.size = Pt(10.5)

    # running header: course code + name
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(hp, 0, 2)
    r = hp.add_run(f"{COURSE_CODE} — {COURSE_NAME}")
    r.bold = True
    r.font.size = Pt(9.5)
    r.font.color.rgb = BRAND
    pPr = hp._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), BRAND_HEX)
    bdr.append(bottom)
    pPr.append(bdr)

    # footer: name / regno  ...  Page X of Y
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tight(fp, 2, 0)
    r = fp.add_run(f"{STUDENT} · {REGNO}     Page ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    field(fp, "PAGE").font.size = Pt(8)
    r = fp.add_run(" of ")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    field(fp, "NUMPAGES").font.size = Pt(8)

    title_page(doc)
    doc.add_page_break()

    b = DocBuilder(doc)
    n_img = n_code = n_out = 0

    for cell in nb.cells:
        if cell.cell_type == "markdown":
            src = cell.source
            if "LABORATORY RECORD" in src:       # replaced by title_page()
                continue
            if src.lstrip().startswith("<div"):  # any other raw HTML block
                continue
            b.markdown(src)

        elif cell.cell_type == "code":
            if cell.source.strip():
                b.code_block(cell.source)
                n_code += 1
            for o in cell.get("outputs", []):
                t = o.get("output_type")
                if t == "stream":
                    b.output_text(o.get("text", ""))
                    n_out += 1
                elif t == "error":
                    b.output_text("\n".join(o.get("traceback", [])), is_error=True)
                elif t in ("execute_result", "display_data"):
                    data = o.get("data", {})
                    if "image/png" in data:
                        b.image(base64.b64decode(data["image/png"]))
                        n_img += 1
                    elif "text/plain" in data:
                        txt = data["text/plain"]
                        if not txt.startswith("<Figure"):
                            b.output_text(txt)
                            n_out += 1

    doc.save(out)
    print(f"{out.name}: {out.stat().st_size/1e6:.1f} MB")
    print(f"  {n_code} code blocks, {n_out} output blocks, {n_img} figures")
    print(f"  header    : {COURSE_CODE} — {COURSE_NAME}")
    print(f"  page border: 2.25pt {BRAND_HEX} box, 24pt from the page edge")
    return 0


if __name__ == "__main__":
    sys.exit(main())
